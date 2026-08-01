import json
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


PLATFORMS = ["kyobo", "yes24", "ridi", "aladin", "millie"]


def iso_date(value):
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    return str(value or "")


def isbn(value):
    digits = "".join(character for character in str(value or "") if character.isdigit())
    return digits


def product(row, headers):
    values = dict(zip(headers, row))
    file_format = str(values.get("파일형태") or "EPUB").upper()
    active_isbn = isbn(values.get("pdfISBN") if file_format == "PDF" else values.get("epubISBN"))
    kyobo_status = "유통중" if values.get("판매상태") == "정상" else "중지"
    platforms = {platform: "확인필요" for platform in PLATFORMS}
    platforms["kyobo"] = kyobo_status
    title = str(values.get("도서명") or "").strip()
    product_code = str(values.get("상품코드") or "").strip()
    return {
        "id": f"kyobo_all_{product_code or values.get('No.')}",
        "assetId": "",
        "recordNo": int(values.get("No.") or 0),
        "title": title,
        "subtitle": "",
        "author": str(values.get("저자명") or "").strip(),
        "publisher": str(values.get("출판사") or "").strip(),
        "category": str(values.get("분류") or "").strip(),
        "price": int(values.get("정가") or 0),
        "fileFormat": file_format,
        "saleStatus": str(values.get("판매상태") or "확인필요"),
        "pdfIsbn": isbn(values.get("pdfISBN")),
        "epubIsbn": isbn(values.get("epubISBN")),
        "activeIsbn": active_isbn,
        "isbnStatus": "발급완료" if active_isbn else "확인필요",
        "publicationDate": iso_date(values.get("출간일")),
        "registeredAt": iso_date(values.get("등록일")),
        "kyoboProductCode": product_code,
        "kyoboSalesProductId": str(values.get("판매상품ID") or "").strip(),
        "salesChannelCode": str(values.get("판매채널코드") or "").strip(),
        "b2bPurchaseStatus": str(values.get("B2B구매상품상태") or "").strip(),
        "b2bRentalStatus": str(values.get("B2B대여상품상태") or "").strip(),
        "samPremiumStatus": str(values.get("SAM프리미엄상품상태") or "").strip(),
        "samUnlimitedStatus": str(values.get("SAM무제한상품상태") or "").strip(),
        "final": True,
        "cover": "",
        "epub": "",
        "pdf": "",
        "manuscript": "",
        "summary": "",
        "platforms": platforms,
        "registeredCount": 1,
        "missingPlatforms": [],
        "attentionPlatforms": ["예스24", "리디", "알라딘", "밀리"],
        "assetMatchMethod": "kyobo-all-time",
        "fivePlatformMatchMethod": "unverified",
        "source": "교보 전체 보유 원장",
    }


def main():
    source = Path(sys.argv[1]).resolve()
    target = Path(sys.argv[2]).resolve()
    sheet = load_workbook(source, read_only=True, data_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    headers = [str(value or "") for value in rows[1]]
    books = []
    for row in rows[2:]:
        if not row or row[0] is None:
            continue
        values = dict(zip(headers, row))
        if str(values.get("파일형태") or "").upper() not in {"PDF", "EPUB"}:
            continue
        books.append(product(row, headers))
    unique_titles = {book["title"].casefold() for book in books}
    payload = {
        "generatedAt": datetime.now().astimezone().isoformat(),
        "sourceWorkbook": str(source),
        "sourceWorkbookUpdatedAt": datetime.fromtimestamp(source.stat().st_mtime).astimezone().isoformat(),
        "fivePlatformWorkbook": "",
        "fivePlatformWorkbookUpdatedAt": "",
        "count": len(books),
        "uniqueBookCount": len(unique_titles),
        "books": books,
    }
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.with_suffix(f".{target.suffix}.{datetime.now().timestamp()}.tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(target)
    print(json.dumps({"count": len(books), "uniqueBookCount": len(unique_titles)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
